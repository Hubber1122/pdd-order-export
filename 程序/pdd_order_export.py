# -*- coding: utf-8 -*-
"""
拼多多买家订单导出 —— 独立可运行程序

用法：
  1) 用配置文件：
       python pdd_order_export.py --config config.json
  2) 用命令行参数：
       python pdd_order_export.py --pdduid <多多号/uid> --token <PDDAccessToken> \
                                  [--type all] [--format xlsx] [--outdir ./output]

说明：只需拼多多登录态凭证（PDDAccessToken + pdduid）即可独立完成一次完整导出，
不依赖原软件任何其他模块。
"""

import os
import sys
import json
import argparse
import datetime
import csv

from field_mapping import FIELD_MAPPING, get_labels
from pdd_client import PddOrderClient, extract_order


# --------------------------------------------------------------------------- #
# 配置加载
# --------------------------------------------------------------------------- #
DEFAULTS = {
    "pdduid": "",
    "access_token": "",
    "verify_auth_token": "",
    "anti_token": "",
    "anti_content": "",
    "order_type": "all",
    "page_size": 10,
    "max_pages": 0,
    "start_offset": "",
    "keyword": "",
    "output_format": "xlsx",     # xlsx | csv | both
    "output_dir": "output",
    "amount_unit": "cent",       # cent=分(默认) | yuan=元
    "time_from": "",
    "time_to": "",
    "need_decrypt": False,
    "fields": [],                # 空=全部字段
}


def load_config(path):
    cfg = dict(DEFAULTS)
    if path and os.path.exists(path):
        with open(path, "r", encoding="utf-8") as f:
            user = json.load(f)
        cfg.update({k: v for k, v in user.items() if k in DEFAULTS})
    return cfg


# --------------------------------------------------------------------------- #
# 字段选择
# --------------------------------------------------------------------------- #
def select_fields(fields):
    if not fields:
        return list(FIELD_MAPPING)
    wanted = set(fields)
    return [f for f in FIELD_MAPPING if f["name"] in wanted]


# --------------------------------------------------------------------------- #
# 时间筛选
# --------------------------------------------------------------------------- #
def parse_time(s):
    if not s:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d %H:%M", "%Y/%m/%d"):
        try:
            return datetime.datetime.strptime(s, fmt)
        except ValueError:
            continue
    raise ValueError(f"无法解析时间: {s}")


def filter_by_time(rows, mapping, time_from, time_to):
    tf = parse_time(time_from)
    tt = parse_time(time_to)
    if not tf and not tt:
        return rows
    order_time_idx = None
    for i, f in enumerate(mapping):
        if f["name"] == "orderTime":
            order_time_idx = i
            break
    out = []
    for row in rows:
        raw = row.get("_raw_order_time") if isinstance(row, dict) else None
        if raw is None:
            # 从已格式化字符串解析
            if order_time_idx is not None:
                val = list(row.values())[order_time_idx] if isinstance(row, dict) else None
                try:
                    raw = datetime.datetime.strptime(str(val), "%Y-%m-%d %H:%M:%S")
                except (ValueError, TypeError):
                    raw = None
        if raw is not None:
            if tf and raw < tf:
                continue
            if tt and raw > tt:
                continue
        out.append(row)
    return out


# --------------------------------------------------------------------------- #
# Excel 导出
# --------------------------------------------------------------------------- #
def write_xlsx(path, rows, mapping):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "多多订单"

    labels = [f["label"] for f in mapping]
    ws.append(labels)

    header_fill = PatternFill("solid", fgColor="C19EE5")
    for col in range(1, len(labels) + 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        values = []
        for f in mapping:
            v = row.get(f["name"], "")
            values.append(v if v is not None else "")
        ws.append(values)

    # 列宽（沿用原实现的 width）
    for i, f in enumerate(mapping, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(6, int(f.get("width", 12)))
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(vertical="top", wrap_text=False)

    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def write_csv(path, rows, mapping):
    labels = [f["label"] for f in mapping]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(labels)
        for row in rows:
            w.writerow([row.get(f["name"], "") for f in mapping])
    return path


# --------------------------------------------------------------------------- #
# 主流程
# --------------------------------------------------------------------------- #
def run(cfg):
    pdduid = cfg["pdduid"].strip()
    token = cfg["access_token"].strip()
    if not pdduid or not token:
        print("[错误] 缺少 pdduid 或 access_token，请在 config.json 中填写，或通过命令行传入。")
        sys.exit(2)

    mapping = select_fields(cfg["fields"])

    client = PddOrderClient(
        pdduid=pdduid,
        access_token=token,
        verify_auth_token=cfg.get("verify_auth_token") or None,
        anti_content=cfg.get("anti_content") or None,
        anti_token=cfg.get("anti_token") or None,
    )

    print(f"[*] 开始拉取订单（type={cfg['order_type']}）...")
    outdir = cfg.get("output_dir") or "output"
    debug_dir = os.path.join(outdir, "debug")
    orders, meta = client.fetch_all(
        order_type=cfg["order_type"],
        start_offset=cfg.get("start_offset") or None,
        max_pages=int(cfg.get("max_pages") or 0),
        page_size=int(cfg.get("page_size") or 10),
        keyword=cfg.get("keyword") or None,
        progress=lambda p, n: print(f"    第 {p} 页，累计 {n} 条"),
        debug_dir=debug_dir,
    )
    print(f"[*] 共获取 {len(orders)} 条订单")

    if not orders:
        print("[!] 未获取到订单（可能是凭证失效或时间范围内无订单）。")
        return

    # 逐条映射为导出行
    need_decrypt = bool(cfg.get("need_decrypt"))
    rows = []
    for o in orders:
        row = extract_order(o, mapping=mapping, need_decrypt=need_decrypt)
        # 记录原始时间戳用于时间筛选
        row["_raw_order_time"] = _raw_time(o)
        rows.append(row)

    # 时间筛选
    rows = filter_by_time(rows, mapping, cfg.get("time_from"), cfg.get("time_to"))
    print(f"[*] 时间筛选后剩余 {len(rows)} 条")

    # 输出
    outdir = cfg.get("output_dir") or "output"
    os.makedirs(outdir, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M")
    base = os.path.join(outdir, f"{pdduid}-多多订单-{stamp}")

    fmt = cfg.get("output_format", "xlsx").lower()
    files = []
    if fmt in ("xlsx", "both"):
        p = write_xlsx(base + ".xlsx", rows, mapping)
        files.append(p)
    if fmt in ("csv", "both"):
        p = write_csv(base + ".csv", rows, mapping)
        files.append(p)

    for p in files:
        print(f"[✓] 已导出: {p}（{len(rows)} 行）")
    return files


def _raw_time(order):
    for k in ("orderTime", "createTime", "order_time", "createdTime"):
        v = order.get(k)
        if v:
            try:
                t = float(v)
                if t > 1e12:
                    t /= 1000.0
                return datetime.datetime.fromtimestamp(t)
            except (ValueError, TypeError, OSError):
                pass
    return None


def main():
    ap = argparse.ArgumentParser(description="拼多多买家订单导出（独立程序）")
    ap.add_argument("--config", default="config.json", help="配置文件路径（默认 config.json）")
    ap.add_argument("--pdduid", help="多多号 / pdd_user_id")
    ap.add_argument("--token", help="PDDAccessToken")
    ap.add_argument("--type", help="订单类型 all/1/2/3/4/5")
    ap.add_argument("--format", help="输出格式 xlsx/csv/both")
    ap.add_argument("--outdir", help="输出目录")
    ap.add_argument("--start-offset", help="起始订单号（断点续传）")
    ap.add_argument("--keyword", help="订单关键词（搜索）")
    ap.add_argument("--from", dest="time_from", help="起始时间 yyyy-MM-dd")
    ap.add_argument("--to", dest="time_to", help="结束时间 yyyy-MM-dd")
    ap.add_argument("--max-pages", type=int, help="最大页数（0=全部）")
    ap.add_argument("--dump", action="store_true", help="导出原始响应 JSON 供调试")
    args = ap.parse_args()

    cfg = load_config(args.config)

    # 命令行覆盖
    overrides = {
        "pdduid": args.pdduid, "access_token": args.token, "order_type": args.type,
        "output_format": args.format, "output_dir": args.outdir,
        "start_offset": args.start_offset, "keyword": args.keyword,
        "time_from": args.time_from, "time_to": args.time_to,
    }
    for k, v in overrides.items():
        if v is not None:
            cfg[k] = v
    if args.max_pages is not None:
        cfg["max_pages"] = args.max_pages

    if args.dump:
        _dump(cfg)
        return

    run(cfg)


def _dump(cfg):
    """调试模式：打印原始响应，帮助校准字段映射。"""
    client = PddOrderClient(cfg["pdduid"].strip(), cfg["access_token"].strip(),
                            verify_auth_token=cfg.get("verify_auth_token") or None,
                            anti_content=cfg.get("anti_content") or None,
                            anti_token=cfg.get("anti_token") or None)
    post_data = {"origin_host_name": "mobile.yangkeduo.com", "page": 1,
                 "size": int(cfg.get("page_size") or 10), "type": cfg.get("order_type", "all")}
    if cfg.get("start_offset"):
        post_data["offset"] = cfg["start_offset"]
    result = client.fetch_page(post_data)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
