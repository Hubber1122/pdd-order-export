# -*- coding: utf-8 -*-
"""
自测脚本：验证字段提取、手机号解密、Excel/CSV 导出逻辑（不依赖真实网络）。
运行：python test_mock.py
"""

import os
import json
import base64
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from field_mapping import FIELD_MAPPING, get_labels
from pdd_client import extract_order, _fmt_amount, _fmt_timestamp
from pdd_decrypt import decrypt_phone, _decrypt_aes
from pdd_order_export import write_xlsx, write_csv


def test_decrypt():
    print("== 测试手机号 AES 解密 ==")
    from Crypto.Cipher import AES
    # 构造与拼多多一致的 secret 派生规则
    # 前 32 位(n)为被索引的字符池，后 32 位(o)为 hex 索引(0-9a-f)
    n = "abcdefghijklmnop" + "qrstuvwxyz012345"  # 32 字符池
    o = "0123456789abcdef" + "0123456789abcdef"  # 32 hex 索引
    secret = n + o
    a = list(n[:16]); l = list(n[16:])
    s = list(o[:16]); u = list(o[16:])
    key = "".join(a[int(c, 16)] for c in s)
    iv = "".join(l[int(c, 16)] for c in u)
    plaintext = "13800138000"
    pad_len = 16 - (len(plaintext) % 16)
    pt = (plaintext + chr(pad_len) * pad_len).encode()
    ct = AES.new(key.encode(), AES.MODE_CBC, iv.encode()).encrypt(pt)
    ciphertext = base64.b64encode(ct).decode()
    dec = decrypt_phone(ciphertext, "0" * 34 + "1" + secret + "|" + "x" * 66)
    assert dec == plaintext, f"解密失败: {dec}"
    print(f"  plaintext={plaintext} -> ciphertext={ciphertext[:20]}... -> 解密={dec}  ✓")


def test_extract():
    print("== 测试字段提取 ==")
    mock_order = {
        "order_sn": "261023-038068590182876",
        "group_order_id": "3478265735434593838",
        "order_amount": 1789,
        "shipping_amount": 0,
        "discount_amount": 0,
        "order_goods": [
            {"goods_id": "123456", "goods_name": "测试商品A", "goods_number": 2,
             "goods_price": 5000, "thumb_url": "https://img/a.jpg", "spec": "红色;M"},
            {"goods_id": "789", "goods_name": "测试商品B", "goods_number": 1,
             "goods_price": 2345, "thumb_url": "https://img/b.jpg", "spec": "蓝色;L"},
        ],
        "order_time": 1792756206,
        "order_status_prompt": "待发货",
        "mall": {"mall_name": "测试旗舰店"},
        "tracking_number": "SF1234567890",
    }
    row = extract_order(mock_order)
    assert row["orderId"] == "261023-038068590182876"
    assert row["productTitle"] == "测试商品A/测试商品B"
    assert row["quantity"] == "3"
    assert row["realPrice"] == "50.00/23.45"
    assert row["actualFee"] == "17.89"
    assert row["tradeStatus"] == "待发货"
    assert row["shopName"] == "测试旗舰店"
    assert row["logisticsNumber"] == "SF1234567890"
    assert row["orderTime"].startswith("2026-10-23")
    print("  orderId   =", row["orderId"], "✓")
    print("  productTitle =", row["productTitle"], "✓")
    print("  quantity  =", row["quantity"], "✓")
    print("  realPrice =", row["realPrice"], "✓")
    print("  actualFee =", row["actualFee"], "✓")
    print("  tradeStatus =", row["tradeStatus"], "✓")
    print("  orderTime =", row["orderTime"], "✓")


def test_export():
    print("== 测试 Excel/CSV 导出 ==")
    os.makedirs("_test_out", exist_ok=True)
    mock_order = {
        "order_sn": "261022-668220823142876",
        "order_goods": [{"goods_name": "测试商品", "goods_number": 1, "goods_price": 9900}],
        "order_amount": 9900,
        "order_time": 1792675758, "order_status_prompt": "待收货",
        "mall": {"mall_name": "测试店"},
    }
    row = extract_order(mock_order)
    rows = [row]
    x = write_xlsx("_test_out/test.xlsx", rows, FIELD_MAPPING)
    c = write_csv("_test_out/test.csv", rows, FIELD_MAPPING)
    assert os.path.exists(x) and os.path.getsize(x) > 0
    assert os.path.exists(c) and os.path.getsize(c) > 0
    print("  xlsx:", x, os.path.getsize(x), "bytes ✓")
    print("  csv :", c, os.path.getsize(c), "bytes ✓")
    # 校验表头
    from openpyxl import load_workbook
    wb = load_workbook(x)
    ws = wb.active
    assert ws.cell(1, 1).value == "订单编号"
    assert ws.cell(2, 1).value == "261022-668220823142876"
    print("  表头/数据校验 ✓ (", ws.max_row, "行 x", ws.max_column, "列 )")


if __name__ == "__main__":
    test_decrypt()
    test_extract()
    test_export()
    print("\n全部自测通过 ✓")
